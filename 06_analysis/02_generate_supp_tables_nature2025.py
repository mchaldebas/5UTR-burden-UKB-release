"""
107.generate_supp_tables_nature2025.py
--------------------------------------
Generates two supplementary tables for the Nature 2025 comparison section.

Table S6        — All 42 Nature 2025 significant 5'UTR associations with
                  5ULTRA detection status and diagnosis. Detection is
                  classified LIVE from Master_Results_Clean.csv.gz using the
                  same concordant Bonferroni/nominal scheme as Results §2 and
                  01_nature2025_overlap.py (n.b. counts are computed, not baked).
Table S7        — Novel 5ULTRA gene-phenotype associations absent from the
                  Consortium's catalog, plus high-gain genes (counts printed
                  at run time; titles set dynamically).

Output: Supp_Tables_S6_S7.xlsx
"""

import os
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side,
                              GradientFill)
from openpyxl.utils import get_column_letter

# ── CONFIG ────────────────────────────────────────────────────────────────────
import sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))
from analysis_config import (DATA_DIR, MASTER_CSV, TRUE_K_MIN,
                             THRESH_U as GWS_ME, MODELS_5ULTRA, MODELS_CADD,
                             N_TESTS_5ULTRA, describe)

NATURE_FILE  = DATA_DIR / 'Nature_2025_Reproduction_Data.csv'
GWS_NAT      = 8.0
K_SUSPICIOUS = 1000
print(describe())

PHENO_LABELS = {
    30000:'WBC count',
    30010:'RBC count',       30020:'Haemoglobin',     30030:'Haematocrit',
    30040:'MCV',             30050:'MCH',              30060:'MCHC',
    30070:'RBC distrib. width', 30080:'Platelet count', 30090:'Platelet crit',
    30100:'Mean platelet vol.', 30110:'Plt distrib width', 30120:'Lymphocyte count',
    30130:'Monocyte count',  30140:'Neutrophil count',30150:'Eosinophil count',
    30160:'Basophil count',  30180:'Lymphocyte %',    30190:'Monocyte %',
    30200:'Neutrophil %',    30210:'Eosinophil %',    30220:'Basophil %',
    30240:'Reticulocyte %',  30250:'Reticulocyte count', 30260:'Mean retic. vol.',
    30270:'Mean sph. cell vol.', 30280:'Immature retic. frac.',
    30290:'High lt. scat. %', 30300:'High lt. scat. count',
    30600:'Albumin',         30610:'Alk. phosphatase',30620:'ALT',
    30630:'Apolipoprotein A',30640:'Apolipoprotein B',30650:'AST',
    30660:'Direct bilirubin',30670:'Urea',            30680:'Calcium',
    30690:'Cholesterol',     30700:'Creatinine',      30710:'CRP',
    30720:'Cystatin C',      30730:'GGT',             30740:'Glucose',
    30750:'HbA1c',
    30760:'HDL cholesterol', 30770:'IGF-1',           30780:'LDL direct',
    30790:'Lipoprotein A',   30800:'Oestradiol',      30810:'Phosphate',
    30820:'Rheumatoid factor',    30830:'SHBG',            30840:'Total bilirubin',
    30850:'Testosterone',    30860:'Total protein',   30870:'Triglycerides',
    30880:'Urate',           30890:'Vitamin D',
}
known_phenos = set(PHENO_LABELS.keys())

# Burden model in the manuscript's vocabulary (Materials and Methods, "Burden
# test design"). Keyed on the MASK, not the model, so the directional model
# resolves to its UP or DN arm rather than collapsing to one label.
# KEEP IN SYNC with MODEL_LABELS in 03_generate_supp_tables_S3_S4.py so every
# supplementary table names the models identically.
MODEL_LABELS = {
    'Mirror_DN':   '5ULTRA DN',
    'Mirror_UP':   '5ULTRA UP',
    '5U_Binary':   '5ULTRA High-confidence',
    '5U_Weighted': '5ULTRA All',
    'Flux_Joint':  '5ULTRA Joint (ACAT)',
    'CD_Binary':   'CADD High-confidence',
    'CD_Weighted': 'CADD All',
}

# ── LOAD DATA ─────────────────────────────────────────────────────────────────
print("Loading data …")
df_all = pd.read_csv(MASTER_CSV)
df_all['gene']  = df_all['gene'].astype(str).str.strip()
df_all['pheno'] = pd.to_numeric(df_all['pheno'], errors='coerce')

u_all   = df_all[df_all['model'].isin(MODELS_5ULTRA)].copy()
u_kpass = u_all[u_all['true_k'] >= TRUE_K_MIN].copy()
u_best_kpass = (u_kpass.sort_values('logp', ascending=False)
                        .drop_duplicates(['gene','pheno']).copy())
u_best_all   = (u_all.sort_values('logp', ascending=False)
                      .drop_duplicates(['gene','pheno']).copy())

u_gws = u_best_kpass[u_best_kpass['logp'] > GWS_ME].copy()
u_gws['pheno_label']   = u_gws['pheno'].map(PHENO_LABELS).fillna(u_gws['pheno'].astype(str))
u_gws['unknown_pheno'] = ~u_gws['pheno'].isin(known_phenos)
u_gws['suspicious_k']  = u_gws['true_k'] > K_SUSPICIOUS

nat_all = pd.read_csv(NATURE_FILE)
nat_all['gene']            = nat_all['gene'].astype(str).str.strip()
nat_all['field_id']        = pd.to_numeric(nat_all['field_id'], errors='coerce')
nat_all['logP_consortium'] = pd.to_numeric(nat_all['-log10P_consortium'], errors='coerce')

# ── LIVE 5ULTRA signal — recomputed from the master CSV (matches 01_nature2025_overlap.py) ──
# FAIR-COMPARISON (replication panel only): Nature imposes no minimum-carrier floor, so we
# use u_best_all (DROP our own k≥5 discovery filter here). Pure numeric join on
# gene + field_id (== master 'pheno'). Direction concordance required for replication.
# The baked '-log10P_me' column is kept only as an audit cross-check (logP_me_baked).
_u_lookup = (u_best_all[['gene', 'pheno', 'logp', 'beta']]
             .rename(columns={'logp': 'logP_me', 'beta': 'beta_me'}))
nat_all = pd.merge(nat_all, _u_lookup,
                   left_on=['gene', 'field_id'], right_on=['gene', 'pheno'],
                   how='left').drop(columns=['pheno'])
nat_all['logP_me']       = nat_all['logP_me'].fillna(0)
nat_all['beta_nat']      = pd.to_numeric(nat_all['beta.5UTR'], errors='coerce')
nat_all['concordant']    = np.sign(nat_all['beta_me']) == np.sign(nat_all['beta_nat'])
nat_all['logP_me_baked'] = pd.to_numeric(nat_all['-log10P_me'], errors='coerce').fillna(0)

# CADD comparison (Aurelie comment #12): best CADD burden −log₁₀P per gene-phenotype,
# same fair-comparison basis as 5ULTRA above (best CADD model, no minimum-carrier
# filter). NaN = the gene-phenotype was never tested by CADD.
c_best_all = (df_all[df_all['model'].isin(MODELS_CADD)]
              .sort_values('logp', ascending=False)
              .drop_duplicates(['gene', 'pheno'])[['gene', 'pheno', 'logp']]
              .rename(columns={'logp': 'logP_cadd'}))
nat_all = pd.merge(nat_all, c_best_all,
                   left_on=['gene', 'field_id'], right_on=['gene', 'pheno'],
                   how='left').drop(columns=['pheno'])

nat = nat_all[nat_all['logP_consortium'] >= GWS_NAT].copy()
nat['pheno_label'] = nat['field_id'].map(PHENO_LABELS).fillna(
    nat['phenotype_name'].astype(str).str.split('#').str[-1].str.strip())

# Replication thresholds — IDENTICAL to the Results §2 / 01_nature2025_overlap.py scheme:
# concordant direction + Bonferroni(0.05/n) primary, nominal(0.05) sensitivity tier.
REPL_NOMINAL = -np.log10(0.05)              # 1.301
REPL_BONF    = -np.log10(0.05 / len(nat))   # ≈2.92 for 42 hits

def classify(row):
    if not bool(row['concordant']):
        return 'Not detected'                 # no 5ULTRA signal, or opposite direction
    if row['logP_me'] > REPL_BONF:
        return 'Detected (Bonferroni)'      # primary
    if row['logP_me'] > REPL_NOMINAL:
        return 'Detected (nominal)'         # sensitivity
    return 'Not detected'

def strength(row):
    # Descriptive sub-label for replicated rows only: 5ULTRA signal strength
    # relative to the consortium value (NOT used to decide replication itself).
    if not str(row['Status']).startswith('Detected'):
        return ''
    lp_me, lp_nat = row['logP_me'], row['logP_consortium']
    if   lp_me >= lp_nat * 1.1: return 'stronger'
    elif lp_me >= lp_nat * 0.8: return 'comparable'
    return 'weaker'

nat['Status']   = nat.apply(classify, axis=1)
nat['Strength'] = nat.apply(strength, axis=1)

# Effect-direction and CADD-comparison display columns (Aurelie comment #12).
def _arrow(b):
    if pd.isna(b) or b == 0:
        return '–'
    return '↑' if b > 0 else '↓'
nat['Direction'] = (nat['beta_nat'].apply(_arrow) + ' / '
                    + nat['beta_me'].apply(_arrow))          # Consortium / 5ULTRA
def _cadd_disp(v):
    if pd.isna(v):
        return 'n/t'          # not tested by CADD
    return f'{v:.1f}' if v > 0 else '< 1'
nat['CADD_disp'] = nat['logP_cadd'].apply(_cadd_disp)

# Missed-hit diagnosis
def diagnose_missed(gene):
    gene_df = u_best_all[u_best_all['gene'] == gene]
    if gene_df.empty:
        return 'Gene absent from 5ULTRA annotation'
    max_k   = gene_df['true_k'].max()
    best_lp = gene_df['logp'].max()
    if max_k < TRUE_K_MIN:
        return f'Insufficient carriers (max k = {max_k:.1f})'
    elif best_lp >= GWS_ME:
        return f'GWS at different phenotype (−log₁₀P = {best_lp:.1f})'
    elif best_lp >= 3.0:
        return f'Suggestive signal (−log₁₀P = {best_lp:.1f})'
    return 'No 5ULTRA signal detected'

def diag(r):
    if str(r['Status']).startswith('Detected'):
        return ''
    # Not detected: distinguish opposite-direction from no-signal
    if pd.notna(r['beta_me']) and r['logP_me'] > 0 and not bool(r['concordant']):
        return f'Opposite direction (5ULTRA −log₁₀P = {r["logP_me"]:.1f})'
    return diagnose_missed(r['gene'])

nat['Notes'] = nat.apply(diag, axis=1)

# ── Replication summary (must match Results §2 / 01_nature2025_overlap.py) ──
_n_bonf = int((nat['Status'] == 'Detected (Bonferroni)').sum())
_n_nom  = int(nat['Status'].isin(['Detected (Bonferroni)', 'Detected (nominal)']).sum())
_n_nd   = int((nat['Status'] == 'Not detected').sum())
print(f"\nConcordance with {len(nat)} Nature 2025 GWS 5'UTR hits (direction-concordant):")
print(f"  Bonferroni (−log₁₀P>{REPL_BONF:.2f}): {_n_bonf}/{len(nat)} ({_n_bonf/len(nat)*100:.0f}%)  [primary]")
print(f"  Nominal    (−log₁₀P>{REPL_NOMINAL:.2f}): {_n_nom}/{len(nat)} ({_n_nom/len(nat)*100:.0f}%)  [sensitivity]")
print(f"  Not detected: {_n_nd}/{len(nat)}")
print(f"  Strength of replicated (Bonf+nominal): "
      f"{nat[nat['Status'].str.startswith('Detected')]['Strength'].value_counts().to_dict()}")

# Sort: Bonferroni-replicated first (by 5ULTRA logP desc), then nominal, then not detected
status_order = {
    'Detected (Bonferroni)': 0,
    'Detected (nominal)':    1,
    'Not detected':            2,
}
nat['_order'] = nat['Status'].map(status_order)
nat_sorted = nat.sort_values(['_order', 'logP_me'], ascending=[True, False]).reset_index(drop=True)

# Novel and high-gain genes
nat_all_genes = set(nat_all['gene'])
novel_all  = u_gws[~u_gws['gene'].isin(nat_all_genes)].drop_duplicates('gene').copy()
novel_clean = novel_all[~novel_all['suspicious_k'] & ~novel_all['unknown_pheno']].copy()
novel_clean = novel_clean.sort_values('logp', ascending=False).reset_index(drop=True)

high_gain = u_gws[u_gws['gene'].isin(nat_all_genes)].copy()
nat_gene_best = (nat_all.groupby('gene')['logP_consortium'].max()
                         .reset_index()
                         .rename(columns={'logP_consortium': 'logP_nat_best'}))
high_gain = pd.merge(high_gain, nat_gene_best, on='gene', how='left')
high_gain['logP_nat_best'] = high_gain['logP_nat_best'].fillna(0)
high_gain = high_gain[high_gain['logP_nat_best'] < GWS_NAT]
high_gain = (high_gain[~high_gain['suspicious_k'] & ~high_gain['unknown_pheno']]
             .sort_values('logp', ascending=False)
             .drop_duplicates('gene')
             .reset_index(drop=True))

print(f"  Novel clean: {len(novel_clean)}  |  High-gain: {len(high_gain)}")

# ── EXCEL HELPERS ─────────────────────────────────────────────────────────────
def make_header_style():
    return {
        'font':      Font(name='Arial', bold=True, size=9, color='FFFFFF'),
        'fill':      PatternFill('solid', fgColor='1F4E79'),
        'alignment': Alignment(horizontal='center', vertical='center',
                               wrap_text=True),
        'border':    Border(
            bottom=Side(style='medium', color='000000'),
            top=Side(style='thin'),
            left=Side(style='thin'),
            right=Side(style='thin'),
        ),
    }

def make_cell_style(bold=False, color=None, center=False):
    fill = PatternFill('solid', fgColor=color) if color else PatternFill()
    return {
        'font':      Font(name='Arial', bold=bold, size=9),
        'fill':      fill,
        'alignment': Alignment(horizontal='center' if center else 'left',
                               vertical='center', wrap_text=True),
        'border':    Border(
            bottom=Side(style='thin', color='CCCCCC'),
            top=Side(style='thin', color='CCCCCC'),
            left=Side(style='thin', color='CCCCCC'),
            right=Side(style='thin', color='CCCCCC'),
        ),
    }

STATUS_COLORS = {
    'Detected (Bonferroni)': 'D6E4F0',
    'Detected (nominal)':    'EBF5FB',
    'Not detected':            'FDFEFE',
}

def apply_style(cell, style):
    for k, v in style.items():
        setattr(cell, k, v)

def write_row(ws, row_idx, values, styles):
    for col_idx, (val, style) in enumerate(zip(values, styles), start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        apply_style(cell, style)

# ── BUILD WORKBOOK ────────────────────────────────────────────────────────────
wb = Workbook()

# ── SHEET 1: Replication table ────────────────────────────────────────────────
ws1 = wb.active
ws1.title = 'S6_Nature_Concordance'
ws1.freeze_panes = 'A3'

# Title row
ws1.merge_cells('A1:I1')
title_cell = ws1['A1']
title_cell.value = ('Supplementary Table S6. 5ULTRA concordance with Nature 2025 '
                    'significant 5\'UTR PheWAS associations (−log₁₀P ≥ 8.0, n = 42)')
title_cell.font      = Font(name='Arial', bold=True, size=10)
title_cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
ws1.row_dimensions[1].height = 28

# Header row
headers_s1 = ['Gene', 'Phenotype', 'Consortium\n−log₁₀P',
               '5ULTRA\n−log₁₀P', 'CADD\n−log₁₀P',
               'Direction\n(Consortium / 5ULTRA)',
               'Status', 'Relative\nstrength',
               'Notes (if not detected)']
hs = make_header_style()
ws1.row_dimensions[2].height = 32
for col, h in enumerate(headers_s1, start=1):
    c = ws1.cell(row=2, column=col, value=h)
    apply_style(c, hs)

# Data rows
for i, row in nat_sorted.iterrows():
    r = i + 3
    status = row['Status']
    bg = STATUS_COLORS.get(status, 'FFFFFF')
    cs = make_cell_style(color=bg)
    cs_c = make_cell_style(color=bg, center=True)
    cs_b = make_cell_style(bold=True, color=bg)

    lp_me_display = f'{row["logP_me"]:.1f}' if row['logP_me'] > 0 else '< 1'

    write_row(ws1, r,
        [row['gene'], row['pheno_label'],
         round(row['logP_consortium'], 1), lp_me_display,
         row['CADD_disp'], row['Direction'],
         status, row['Strength'], row['Notes']],
        [cs_b, cs, cs_c, cs_c, cs_c, cs_c, cs_c, cs_c, cs])

    ws1.row_dimensions[r].height = 16

# Column widths
for col_idx, width in enumerate([10, 22, 14, 14, 12, 20, 22, 14, 40], start=1):
    ws1.column_dimensions[get_column_letter(col_idx)].width = width

# Legend
legend_row = len(nat_sorted) + 4
ws1.cell(row=legend_row, column=1,
         value='Colour key:').font = Font(name='Arial', bold=True, size=8)
for offset, (status, color) in enumerate(STATUS_COLORS.items(), start=1):
    c = ws1.cell(row=legend_row + offset, column=1, value=status)
    c.font = Font(name='Arial', size=8)
    c.fill = PatternFill('solid', fgColor=color)
ws1.cell(row=legend_row + len(STATUS_COLORS) + 1, column=1,
         value=(f'Concordance (fair comparison, no minimum-carrier filter; direction-concordant): '
                f'Bonferroni −log₁₀P > {REPL_BONF:.2f} (0.05/{len(nat)}) primary; '
                f'nominal −log₁₀P > {REPL_NOMINAL:.2f} (P < 0.05) sensitivity. '
                f'"Relative strength" compares the 5ULTRA −log₁₀P to the consortium value '
                f'(stronger ≥1.1×, comparable 0.8–1.1×, weaker <0.8×); it is descriptive only '
                f'and does not affect detection status. "Direction" gives the sign of the '
                f'burden effect for the Consortium and for 5ULTRA (↑ positive, ↓ negative, – no '
                f'signal); detection requires concordant direction. "CADD −log₁₀P" is the best '
                f'CADD burden result for the same gene-phenotype (same no-carrier-filter basis; '
                f'n/t = not tested by CADD), shown for comparison.')).font = Font(name='Arial', italic=True, size=8)

# ── SHEET 2: Novel + High-gain ────────────────────────────────────────────────
ws2 = wb.create_sheet('S7_Novel_Genes')
ws2.freeze_panes = 'A3'

# Title
ws2.merge_cells('A1:G1')
t2 = ws2['A1']
t2.value = (f'Supplementary Table S7. Novel 5ULTRA gene-phenotype associations '
            f'absent from the Nature 2025 Consortium catalog (n = {len(novel_clean)}), '
            f'and genes with substantially higher 5ULTRA power (n = {len(high_gain)})')
t2.font      = Font(name='Arial', bold=True, size=10)
t2.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
ws2.row_dimensions[1].height = 36

# Header
headers_s2 = ['Gene', 'Phenotype', '5ULTRA\n−log₁₀P',
               'Effect size\n(β, SD units)', 'Weighted burden\nsum (k)',
               'Best model', 'Category']
ws2.row_dimensions[2].height = 32
for col, h in enumerate(headers_s2, start=1):
    c = ws2.cell(row=2, column=col, value=h)
    apply_style(c, hs)

# Novel genes section header
NOVEL_COLOR  = 'D5E8D4'   # green tint
HIGAIN_COLOR = 'DAE8FC'   # blue tint

section_row = 3
c = ws2.cell(row=section_row, column=1,
             value='Part A — Novel genes: entirely absent from the Nature 2025 catalog')
c.font      = Font(name='Arial', bold=True, size=9, color='1D6320')
c.fill      = PatternFill('solid', fgColor='A9D18E')
c.alignment = Alignment(horizontal='left', vertical='center')
ws2.merge_cells(f'A{section_row}:G{section_row}')
ws2.row_dimensions[section_row].height = 16

for i, row in novel_clean.iterrows():
    r = section_row + 1 + i
    cs   = make_cell_style(color=NOVEL_COLOR)
    cs_c = make_cell_style(color=NOVEL_COLOR, center=True)
    cs_b = make_cell_style(bold=True, color=NOVEL_COLOR)
    write_row(ws2, r,
        [row['gene'], row['pheno_label'],
         round(row['logp'], 1), round(row['beta'], 3),
         round(float(row['true_k']), 1), MODEL_LABELS.get(row['mask'], row['mask']),
         'Novel'],
        [cs_b, cs, cs_c, cs_c, cs_c, cs, cs_c])
    ws2.row_dimensions[r].height = 16

# High-gain section header
hg_section_row = section_row + 1 + len(novel_clean) + 1
c2 = ws2.cell(row=hg_section_row, column=1,
              value='Part B — Higher-power genes: present in Nature 2025 catalog but sub-threshold (−log₁₀P < 8)')
c2.font      = Font(name='Arial', bold=True, size=9, color='1A3E6E')
c2.fill      = PatternFill('solid', fgColor='9DC3E6')
c2.alignment = Alignment(horizontal='left', vertical='center')
ws2.merge_cells(f'A{hg_section_row}:G{hg_section_row}')
ws2.row_dimensions[hg_section_row].height = 16

# Extra header for high-gain (replace "Effect size" with "Nature 2025 best −log10P")
for r_off, row in high_gain.iterrows():
    r = hg_section_row + 1 + r_off
    cs   = make_cell_style(color=HIGAIN_COLOR)
    cs_c = make_cell_style(color=HIGAIN_COLOR, center=True)
    cs_b = make_cell_style(bold=True, color=HIGAIN_COLOR)
    nat_best_str = f'{row["logP_nat_best"]:.1f}'
    write_row(ws2, r,
        [row['gene'], row['pheno_label'],
         round(row['logp'], 1), nat_best_str,
         round(float(row['true_k']), 1), MODEL_LABELS.get(row['mask'], row['mask']),
         'Higher power'],
        [cs_b, cs, cs_c, cs_c, cs_c, cs, cs_c])
    ws2.row_dimensions[r].height = 16

# Column B note for high-gain section
note_row = hg_section_row + 1 + len(high_gain) + 1
ws2.cell(row=note_row, column=1,
         value='* For Part B, the "Effect size" column shows the Nature 2025 consortium best −log₁₀P for the gene (across all phenotypes tested).').font = Font(name='Arial', italic=True, size=8)
ws2.merge_cells(f'A{note_row}:G{note_row}')

ws2.cell(row=note_row + 1, column=1,
         value=f'5ULTRA GWS threshold: −log₁₀P > {GWS_ME:.3f} (Bonferroni correction for the {N_TESTS_5ULTRA:,} association tests performed in the 5ULTRA suite at k ≥ {TRUE_K_MIN:.0f}: 0.05/{N_TESTS_5ULTRA:,}). k is the aggregate weighted allele count for the best model × spectrum (the sum of weighted allele dosages across all individuals, REGENIE --build-mask sum), not a raw carrier count; pairs with k < {TRUE_K_MIN:.0f} were excluded.').font = Font(name='Arial', italic=True, size=8)
ws2.merge_cells(f'A{note_row+1}:G{note_row+1}')

# Column widths
for col_idx, width in enumerate([10, 22, 13, 14, 14, 18, 14], start=1):
    ws2.column_dimensions[get_column_letter(col_idx)].width = width

# ── SAVE ─────────────────────────────────────────────────────────────────────
out = 'Supp_Tables_S6_S7.xlsx'
wb.save(out)
print(f'Saved: {out}')
print(f'  Sheet 1: {len(nat_sorted)} rows (42 Nature 2025 associations)')
print(f'  Sheet 2: {len(novel_clean)} novel + {len(high_gain)} high-gain genes')
