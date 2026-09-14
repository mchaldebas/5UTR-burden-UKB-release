"""
Generate Supplementary Table S2 — sensitivity of the 5'UTR-DN vs exome-PTV
correlation to the carrier-count floor.

Answers Aurelie's question about where the 23 pairs come from: the table shows
the full funnel and demonstrates that the correlation is significant at every
floor, so k >= 15 is a presentational choice for Figure 4b rather than one that
generates the result.

Input:  ptv_k_floor_sensitivity.tsv  (from 13_ptv_k_floor_sensitivity.py)
Output: Supp_Table_S2.xlsx

Styling matches 03_generate_supp_tables_S3_S4.py so S1-S6 look like one set.

Usage:
    python 06_analysis/13_ptv_k_floor_sensitivity.py     # writes the TSV
    python 06_analysis/14_generate_supp_table_S2.py
"""

import os
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

IN_TSV   = Path(os.environ.get("PTV_SENS_TSV", "ptv_k_floor_sensitivity.tsv"))
OUT_XLSX = "Supp_Table_S2.xlsx"

# Palette lifted from 03_generate_supp_tables_S3_S4.py
C_HEADER  = "BDD7EE"
C_STRIPE  = "EAF4FB"
C_PRIMARY = "E8F8E8"   # highlight the floor used in the manuscript
C_WHITE   = "FFFFFF"
C_BORDER  = "BFBFBF"
PRIMARY_FLOOR = 15

if not IN_TSV.exists():
    sys.exit(f"ERROR: {IN_TSV} not found — run 13_ptv_k_floor_sensitivity.py first.")

d = pd.read_csv(IN_TSV, sep="\t")

# The UP negative control columns are not reproduced reliably by script 13
# (its model-name filter does not match the UP rows in the triad file); the
# manuscript takes that control from figure4.py. Drop them rather than ship
# empty columns.
d = d.drop(columns=[c for c in d.columns if "UP_control" in c], errors="ignore")

COLS = [
    ("k_floor",         "Carrier floor (k ≥)", 18),
    ("n_pairs",         "Gene-phenotype pairs", 20),
    ("pearson_r",       "Pearson r",            12),
    ("pearson_p",       "Pearson P",            14),
    ("spearman_rho",    "Spearman ρ",           13),
    ("spearman_p",      "Spearman P",           14),
    ("n_concordant",    "Concordant direction", 20),
    ("pct_concordant",  "% concordant",         14),
    ("n_with_missense", "With missense mask",   19),
]
COLS = [(k, lab, w) for k, lab, w in COLS if k in d.columns]

def thin():
    s = Side(style="thin", color=C_BORDER)
    return Border(left=s, right=s, top=s, bottom=s)

wb = Workbook()
ws = wb.active
ws.title = "S2 - PTV correlation vs k"

ncol = len(COLS)

# Title
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
t = ws.cell(row=1, column=1, value=(
    "Supplementary Table S2. Sensitivity of the 5′UTR-DN versus exome PTV "
    "correlation to the carrier-count floor. Each row repeats the analysis of "
    "Figure 4a at a different minimum 5′UTR-DN carrier count. The row used in "
    f"the manuscript (k ≥ {PRIMARY_FLOOR}) is shaded."))
t.font = Font(name="Calibri", bold=True, size=11)
t.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
t.fill = PatternFill("solid", fgColor=C_HEADER)
ws.row_dimensions[1].height = 30

# Note
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
n = ws.cell(row=2, column=1, value=(
    "Pairs are those with an unweighted 5ULTRA DN burden (af5, score ≥ 0.5) and "
    "a published exome protein-truncating variant statistic from the UK Biobank "
    "WGS consortium, matched on gene and phenotype; they are not restricted to "
    "the genome-wide significant 5ULTRA associations. 'With missense mask' counts "
    "pairs whose best consortium coding model was a damaging-missense or "
    "nonsynonymous mask, the only pairs for which a missense effect is shown in "
    "Figure 4b."))
n.font = Font(name="Calibri", italic=True, size=9, color="595959")
n.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws.row_dimensions[2].height = 42

# Header
hdr = 3
for j, (_, label, width) in enumerate(COLS, start=1):
    c = ws.cell(row=hdr, column=j, value=label)
    c.font = Font(name="Calibri", bold=True, size=10)
    c.fill = PatternFill("solid", fgColor=C_HEADER)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = thin()
    ws.column_dimensions[get_column_letter(j)].width = width
ws.row_dimensions[hdr].height = 30

# Body
for i, (_, row) in enumerate(d.iterrows()):
    r = hdr + 1 + i
    is_primary = int(row.get("k_floor", -1)) == PRIMARY_FLOOR
    fill = C_PRIMARY if is_primary else (C_STRIPE if i % 2 else C_WHITE)
    for j, (key, _, _) in enumerate(COLS, start=1):
        v = row[key]
        c = ws.cell(row=r, column=j, value=(0 if key == "k_floor" and v == 0 else v))
        c.font = Font(name="Calibri", size=10, bold=is_primary)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin()

ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
wb.save(OUT_XLSX)
print(f"Saved: {OUT_XLSX}  ({len(d)} carrier floors)")
