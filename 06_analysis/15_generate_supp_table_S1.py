"""
Generate Supplementary Table S1 — phenotype abbreviations and full names.

Aurelie comment #15: figures and tables use abbreviated trait labels
("Plt distrib width", "High lt scat %"), which are not self-explanatory. This
table gives, for each of the 59 traits, the UK Biobank field ID, the
abbreviation used in the figures, and the full phenotype name.

The abbreviation column is READ FROM pheno_labels.py rather than retyped, so
the table cannot drift from the labels actually drawn on the figures. If a
figure overlays its own shorter form (figure3.py does, for the dense forest-plot
rows), that variant is shown in a separate column.

Output: Supp_Table_S1.xlsx
Styling matches 03_generate_supp_tables_S3_S4.py so S1-S7 look like one set.

Usage:
    python 06_analysis/15_generate_supp_table_S1.py
"""

import re
import sys
import pathlib

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "07_figures"))
from pheno_labels import PHENO_LABELS          # canonical abbreviations

OUT_XLSX = "Supp_Table_S1.xlsx"

C_HEADER = "BDD7EE"
C_STRIPE = "EAF4FB"
C_BIOCH  = "FEF9E7"
C_WHITE  = "FFFFFF"
C_BORDER = "BFBFBF"

# UK Biobank Showcase field names. Blood count = category 100081,
# blood biochemistry = category 17518.
FULL_NAMES = {
    30000: "White blood cell (leukocyte) count",
    30010: "Red blood cell (erythrocyte) count",
    30020: "Haemoglobin concentration",
    30030: "Haematocrit percentage",
    30040: "Mean corpuscular volume",
    30050: "Mean corpuscular haemoglobin",
    30060: "Mean corpuscular haemoglobin concentration",
    30070: "Red blood cell (erythrocyte) distribution width",
    30080: "Platelet count",
    30090: "Platelet crit",
    30100: "Mean platelet (thrombocyte) volume",
    30110: "Platelet distribution width",
    30120: "Lymphocyte count",
    30130: "Monocyte count",
    30140: "Neutrophil count",
    30150: "Eosinophil count",
    30160: "Basophil count",
    30180: "Lymphocyte percentage",
    30190: "Monocyte percentage",
    30200: "Neutrophil percentage",
    30210: "Eosinophil percentage",
    30220: "Basophil percentage",
    30240: "Reticulocyte percentage",
    30250: "Reticulocyte count",
    30260: "Mean reticulocyte volume",
    30270: "Mean sphered cell volume",
    30280: "Immature reticulocyte fraction",
    30290: "High light scatter reticulocyte percentage",
    30300: "High light scatter reticulocyte count",
    30600: "Albumin",
    30610: "Alkaline phosphatase",
    30620: "Alanine aminotransferase",
    30630: "Apolipoprotein A",
    30640: "Apolipoprotein B",
    30650: "Aspartate aminotransferase",
    30660: "Direct bilirubin",
    30670: "Urea",
    30680: "Calcium",
    30690: "Cholesterol",
    30700: "Creatinine",
    30710: "C-reactive protein",
    30720: "Cystatin C",
    30730: "Gamma glutamyltransferase",
    30740: "Glucose",
    30750: "Glycated haemoglobin (HbA1c)",
    30760: "HDL cholesterol",
    30770: "Insulin-like growth factor 1 (IGF-1)",
    30780: "LDL direct",
    30790: "Lipoprotein A",
    30800: "Oestradiol",
    30810: "Phosphate",
    30820: "Rheumatoid factor",
    30830: "Sex hormone-binding globulin (SHBG)",
    30840: "Total bilirubin",
    30850: "Testosterone",
    30860: "Total protein",
    30870: "Triglycerides",
    30880: "Urate",
    30890: "Vitamin D",
}


def figure3_short_forms():
    """Pull figure3.py's _PHENO_SHORT overlay without importing the module
    (importing would try to load the master results file)."""
    src = (ROOT / "07_figures" / "figure3.py").read_text(encoding="utf-8")
    m = re.search(r"_PHENO_SHORT\s*=\s*\{(.*?)\n\}", src, re.S)
    if not m:
        return {}
    return {int(k): v for k, v in re.findall(r"(\d{5})\s*:\s*'([^']*)'", m.group(1))}


SHORT3 = figure3_short_forms()

missing = sorted(set(PHENO_LABELS) - set(FULL_NAMES))
if missing:
    sys.exit(f"ERROR: no full name for field(s) {missing}")

rows = []
for fid in sorted(PHENO_LABELS):
    abbrev = PHENO_LABELS[fid]
    alt = SHORT3.get(fid, "")
    rows.append({
        "fid": fid,
        "abbrev": abbrev,
        "alt": "" if alt in ("", abbrev) else alt,
        "full": FULL_NAMES[fid],
        "cat": "Blood count" if fid < 30500 else "Blood biochemistry",
    })

COLS = [
    ("fid",    "UK Biobank field ID",            20),
    ("abbrev", "Abbreviation used in figures",   30),
    ("alt",    "Alternative short form",         26),
    ("full",   "Full phenotype name",            48),
    ("cat",    "Category",                       20),
]


def thin():
    s = Side(style="thin", color=C_BORDER)
    return Border(left=s, right=s, top=s, bottom=s)


wb = Workbook()
ws = wb.active
ws.title = "S1 - Phenotype abbreviations"
ncol = len(COLS)

ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
t = ws.cell(row=1, column=1, value=(
    "Supplementary Table S1. Abbreviations and full names for the 59 quantitative "
    "traits analysed, with their UK Biobank field IDs. Abbreviations are those "
    "drawn on the figures; where a figure uses a shorter form for space, it is "
    "given in the third column."))
t.font = Font(name="Calibri", bold=True, size=11)
t.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
t.fill = PatternFill("solid", fgColor=C_HEADER)
ws.row_dimensions[1].height = 30

hdr = 2
for j, (_, label, width) in enumerate(COLS, start=1):
    c = ws.cell(row=hdr, column=j, value=label)
    c.font = Font(name="Calibri", bold=True, size=10)
    c.fill = PatternFill("solid", fgColor=C_HEADER)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = thin()
    ws.column_dimensions[get_column_letter(j)].width = width
ws.row_dimensions[hdr].height = 28

for i, row in enumerate(rows):
    r = hdr + 1 + i
    fill = C_BIOCH if row["cat"] == "Blood biochemistry" else (C_STRIPE if i % 2 else C_WHITE)
    for j, (key, _, _) in enumerate(COLS, start=1):
        c = ws.cell(row=r, column=j, value=row[key])
        c.font = Font(name="Calibri", size=10)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center" if j in (1, 5) else "left",
                                vertical="center")
        c.border = thin()

ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
wb.save(OUT_XLSX)

n_bc = sum(r["cat"] == "Blood count" for r in rows)
print(f"Saved: {OUT_XLSX}  ({len(rows)} traits: {n_bc} blood count, "
      f"{len(rows) - n_bc} blood biochemistry)")
if SHORT3:
    print(f"  {sum(1 for r in rows if r['alt'])} traits carry a shorter Figure 3 form")
