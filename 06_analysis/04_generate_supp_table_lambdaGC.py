"""
Generate the Supplementary Table of genomic inflation factors (λGC).

One row per model × allele-frequency spectrum, reporting the number of
association tests and the genomic inflation factor λGC.

Rationale
---------
λGC must be read off a SINGLE fixed model, not the best-of-N per gene-phenotype
pair: taking the max −log₁₀P over several models × spectra is a selection on the
most significant result and inflates λGC by construction, making it
uninterpretable as a calibration metric. This table reports λGC for every main
model so the calibration is fully transparent and the figure's single-model
choice is justified.

The Supplementary Figure 1 QQ panels use the best-calibrated single models,
flagged here with "In Fig. S1":
    5ULTRA Binary (af5)  λGC ≈ 1.01
    CADD   Binary (af5)  λGC ≈ 1.05

Output: Supp_Table_lambdaGC.xlsx

λGC is computed exactly as in 07_figures/supp_figure1.py:
    chi2 = chi2.isf(P, df=1);  λGC = median(chi2) / 0.4549364
with the same true_k ≥ 5 filter applied first.
"""

import os
from pathlib import Path
import numpy as np
import pandas as pd
import scipy.stats as stats
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Constants ──────────────────────────────────────────────────────────────────
DATA_DIR    = Path(os.environ.get('DATA_DIR', '/Volumes/MCHALDEBAS3/UKB-500k/UKB-data'))
MASTER_CSV  = DATA_DIR / 'Master_Results_Clean.csv.gz'
OUT_XLSX    = 'Supp_Table_lambdaGC.xlsx'
TRUE_K_MIN  = 5
CHI2_MEDIAN = 0.4549364   # median of chi-squared(df=1) = qchisq(0.5, 1)

# Main models reported (controls/ablations excluded). Ordered by scorer then
# by mask, so the table reads 5ULTRA block first, then CADD block.
MODEL_ORDER = [
    '5ULTRA_Binary', '5ULTRA_Weighted', '5U_Logic', 'Flux_Joint',
    'CADD_Binary', 'CADD_Weighted',
]
SPEC_ORDER  = ['af5', 'rare']

SCORER_LABELS = {
    '5ULTRA_Binary': '5ULTRA', '5ULTRA_Weighted': '5ULTRA',
    '5U_Logic': '5ULTRA',      'Flux_Joint': '5ULTRA',
    'CADD_Binary': 'CADD',     'CADD_Weighted': 'CADD',
}
MODEL_LABELS = {
    '5U_Logic':        'Logic',
    '5ULTRA_Binary':   'Binary',
    '5ULTRA_Weighted': 'Weighted',
    'Flux_Joint':      'Flux Joint',
    'CADD_Binary':     'Binary',
    'CADD_Weighted':   'Weighted',
}
SPEC_LABELS = {'rare': 'AF < 1%', 'af5': 'AF < 5%'}

# Model × spectrum combinations shown in the Supp. Fig. 1 QQ panels.
FIG_S1_CHOICE = {('5ULTRA_Binary', 'af5'), ('CADD_Binary', 'af5')}

# ── Colour palette (consistent with Supp_Tables_S3_S4) ──────────────────────────
C_HEADER     = 'BDD7EE'   # light blue header
C_ROW_5U     = 'EAF4FB'   # very light blue – 5ULTRA rows
C_ROW_CADD   = 'F2F2F2'   # very light grey – CADD rows
C_ROW_FIG    = 'E8F8E8'   # very light green – row used in the figure
C_BORDER     = 'BFBFBF'

def thin_border():
    s = Side(style='thin', color=C_BORDER)
    return Border(left=s, right=s, top=s, bottom=s)

# ── λGC helper ──────────────────────────────────────────────────────────────────
def lambda_gc(logp):
    logp = np.asarray(logp, dtype=float)
    chi2 = stats.chi2.isf(np.clip(10 ** -logp, 1e-300, 1.0), df=1)
    return float(np.median(chi2) / CHI2_MEDIAN)

# ── Load data ──────────────────────────────────────────────────────────────────
print("Loading Master_Results_Clean.csv.gz …")
df = pd.read_csv(MASTER_CSV, usecols=['model', 'spectrum', 'pheno', 'logp', 'true_k'])
df = df[df['true_k'] >= TRUE_K_MIN]

# Minimum gene tests per trait to compute a meaningful per-trait λGC.
MIN_TRAIT_TESTS = 50

rows = []
for model in MODEL_ORDER:
    for spec in SPEC_ORDER:
        sub = df[(df['model'] == model) & (df['spectrum'] == spec)]
        if sub.empty:
            continue
        lam = lambda_gc(sub['logp'].values)
        # Per-trait λGC (across genes, within each phenotype): shows the pooled
        # value is not driven by correlated phenotypes (Aurelie comment on
        # computing λGC separately by trait).
        per_trait = (sub.groupby('pheno')['logp']
                     .apply(lambda s: lambda_gc(s.values) if len(s) >= MIN_TRAIT_TESTS
                            else np.nan)
                     .dropna())
        if len(per_trait):
            t_med = float(per_trait.median())
            t_lo, t_hi = float(per_trait.min()), float(per_trait.max())
            n_traits = len(per_trait)
            trait_median = round(t_med, 3)
            trait_range = f'{t_lo:.2f}–{t_hi:.2f}'
        else:
            trait_median, trait_range, n_traits = np.nan, '', 0
        rows.append({
            'Scorer':   SCORER_LABELS[model],
            'Model':    MODEL_LABELS[model],
            'Spectrum': SPEC_LABELS[spec],
            'N tests':  len(sub),
            'lambdaGC': round(lam, 3),
            'N traits':      n_traits,
            'trait_median':  trait_median,
            'trait_range':   trait_range,
            'In Fig. S1': 'Yes' if (model, spec) in FIG_S1_CHOICE else '',
            '_scorer_key': SCORER_LABELS[model],
            '_in_fig':     (model, spec) in FIG_S1_CHOICE,
        })
        print(f"  {model:<16} {spec:<5}  N={len(sub):>7,}  λGC={lam:.3f}  "
              f"per-trait median={trait_median} range={trait_range} ({n_traits} traits)")

table = pd.DataFrame(rows)

# ── Write styled Excel sheet ────────────────────────────────────────────────────
COLS       = ['Scorer', 'Model', 'Spectrum', 'N tests', 'λGC',
              'N traits', 'Per-trait λGC (median)', 'Per-trait λGC (range)',
              'In Fig. S1']
COL_WIDTHS = [12, 14, 12, 12, 10, 10, 18, 18, 12]

wb = Workbook()
ws = wb.active
ws.title = 'Genomic inflation (λGC)'

# Title
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLS))
title = ws.cell(row=1, column=1,
                value='Supplementary Table — Genomic inflation factor (λGC) by model and '
                      'allele-frequency spectrum')
title.font      = Font(name='Calibri', bold=True, size=11)
title.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
title.fill      = PatternFill('solid', fgColor=C_HEADER)
ws.row_dimensions[1].height = 28

# Note
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(COLS))
note = ws.cell(row=2, column=1,
    value=f'λGC = median(χ²₁) / {CHI2_MEDIAN:.4f}, computed per model × spectrum on all '
          f'gene-phenotype association tests passing true_k ≥ {TRUE_K_MIN}. Each model × '
          f'spectrum is a single fixed model, so λGC is interpretable as a calibration '
          f'metric (no best-of-N selection across models, which would inflate λGC by '
          f'construction). Per-trait λGC (median and range) is computed separately within '
          f'each phenotype across genes (traits with ≥ {MIN_TRAIT_TESTS} gene tests); its '
          f'closeness to the pooled λGC shows the pooled value is not driven by correlated '
          f'phenotypes. Rows flagged "In Fig. S1" are the best-calibrated single models '
          f'shown in the Supplementary Figure 1 QQ panels.')
note.font      = Font(name='Calibri', italic=True, size=9, color='595959')
note.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
ws.row_dimensions[2].height = 72

# Header row
for ci, col in enumerate(COLS, start=1):
    cell = ws.cell(row=3, column=ci, value=col)
    cell.font      = Font(name='Calibri', bold=True, size=10)
    cell.fill      = PatternFill('solid', fgColor=C_HEADER)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border    = thin_border()
ws.row_dimensions[3].height = 20

# Data rows
for ri, r in enumerate(table.to_dict('records'), start=4):
    if r['_in_fig']:
        fill = PatternFill('solid', fgColor=C_ROW_FIG)
    elif r['_scorer_key'] == '5ULTRA':
        fill = PatternFill('solid', fgColor=C_ROW_5U)
    else:
        fill = PatternFill('solid', fgColor=C_ROW_CADD)

    values = [r['Scorer'], r['Model'], r['Spectrum'], r['N tests'],
              r['lambdaGC'], r['N traits'], r['trait_median'], r['trait_range'],
              r['In Fig. S1']]
    for ci, val in enumerate(values, start=1):
        cell = ws.cell(row=ri, column=ci, value=val)
        cell.fill      = fill
        cell.border    = thin_border()
        cell.alignment = Alignment(horizontal='center' if ci >= 3 else 'left',
                                   vertical='center')
        cell.font      = Font(name='Calibri', size=10, bold=r['_in_fig'])
        if ci == 4:                       # N tests: thousands separator
            cell.number_format = '#,##0'
        elif ci in (5, 7):                # λGC / per-trait median: 3 decimals
            cell.number_format = '0.000'
        elif ci == 6:                     # N traits: integer
            cell.number_format = '0'

# Column widths
for ci, w in enumerate(COL_WIDTHS, start=1):
    ws.column_dimensions[get_column_letter(ci)].width = w

ws.freeze_panes   = 'A4'
ws.auto_filter.ref = f'A3:{get_column_letter(len(COLS))}3'

wb.save(OUT_XLSX)
print(f"\nSaved: {OUT_XLSX}  ({len(table)} rows)")
